"""Document capabilities (v3.1 §14).

Structured operations around the existing deterministic builders
(``agents/documents.py``: pptx/docx/xlsx/pdf):

    document_validate -> document_create -> document_persist -> document_inspect
                                             \\
                                              document_extract / document_convert

Artifacts are first-class: ``document_persist`` records an ``Artifact`` row and
emits ``artifact.created``, so generated files appear in the outputs UI and in
the event trace instead of only existing on disk.

Honest limits: extraction supports text formats, docx and xlsx. PDF extraction
is not implemented (there is no PDF text extractor in the dependency set) and
``document_extract`` says so rather than returning garbage.
"""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path
from typing import Any

from agent_system.services.tool_errors import ToolError
from agent_system.services.tools.builtin._exec import read_text
from agent_system.services.tools.paths import allowed_roots, jailed, relative_label
from agent_system.services.tools.registry import Tool, ToolContext, ToolRegistry, _str_param

GROUP = "documents"
KINDS = ("pptx", "docx", "xlsx", "pdf", "txt", "md")
MAX_TEXT_CHARS = 20000

#: Required content keys per builder (kept in sync with DocumentAgent).
REQUIRED_CONTENT: dict[str, tuple[tuple[str, ...], ...]] = {
    "pptx": (("title",),),
    "docx": (("title",),),
    "xlsx": (("sheets", "rows", "headers"),),
    "pdf": (("title",),),
}


def _outputs_dir(ctx: ToolContext) -> Path:
    roots = allowed_roots(ctx.settings)
    for root in roots:
        if root.name == "outputs":
            root.mkdir(parents=True, exist_ok=True)
            return root
    target = Path.cwd() / "outputs"
    target.mkdir(parents=True, exist_ok=True)
    return target


def _document_agent(ctx: ToolContext) -> Any:
    from agent_system.agents.documents import DocumentAgent

    configured = getattr(ctx.settings, "outputs_dir", None)
    return DocumentAgent(Path(configured) if configured else _outputs_dir(ctx))


def _validate_content(kind: str, content: dict[str, Any]) -> list[str]:
    problems: list[str] = []
    if kind in REQUIRED_CONTENT:
        for alternatives in REQUIRED_CONTENT[kind]:
            if not any(key in content for key in alternatives):
                problems.append(f"{kind} content requires one of: {', '.join(alternatives)}")
    if "title" in content and not isinstance(content["title"], str):
        problems.append("'title' must be a string")
    if "sections" in content and not isinstance(content["sections"], list):
        problems.append("'sections' must be a list")
    return problems


def _content(args: dict[str, Any]) -> dict[str, Any]:
    raw = args.get("content")
    if raw is None:
        return {}
    if isinstance(raw, str):
        import json

        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError as exc:
            raise ToolError(f"document content is not valid JSON: {exc.msg}") from exc
        if not isinstance(parsed, dict):
            raise ToolError("document content must be a JSON object")
        return parsed
    if isinstance(raw, dict):
        return raw
    raise ToolError("document content must be an object")


def _kind(args: dict[str, Any]) -> str:
    kind = str(args.get("kind") or "").lower().lstrip(".")
    if kind not in KINDS:
        raise ToolError(f"unsupported document kind '{kind}' (expected one of: {', '.join(KINDS)})")
    return kind


def _target_name(args: dict[str, Any], kind: str) -> str:
    name = str(args.get("name") or "document").strip() or "document"
    if "/" in name or "\\" in name or name.startswith("."):
        raise ToolError("document name must be a plain file name")
    return name if name.endswith(f".{kind}") else f"{name}.{kind}"


def _document_validate(args: dict[str, Any], ctx: ToolContext) -> dict[str, Any]:  # noqa: ARG001
    kind = _kind(args)
    content = _content(args)
    problems = _validate_content(kind, content)
    return {
        "kind": kind,
        "valid": not problems,
        "problems": problems,
        "keys": sorted(content),
    }


def _document_create(args: dict[str, Any], ctx: ToolContext) -> dict[str, Any]:
    kind = _kind(args)
    content = _content(args)
    problems = _validate_content(kind, content)
    if problems:
        raise ToolError("invalid document content: " + "; ".join(problems))
    name = _target_name(args, kind)
    agent = _document_agent(ctx)
    if kind in ("txt", "md"):
        target = _outputs_dir(ctx) / name
        body = str(content.get("body") or content.get("text") or "")
        if not body.strip():
            raise ToolError(f"{kind} documents require 'body' text")
        target.write_text(body, encoding="utf-8")
    else:
        try:
            target = agent.generate(kind, name, content)
        except Exception as exc:
            raise ToolError(f"document generation failed: {exc}") from exc
    return {
        "kind": kind,
        "path": str(target),
        "name": name,
        "bytes": target.stat().st_size if target.exists() else 0,
    }


def _document_inspect(args: dict[str, Any], ctx: ToolContext) -> dict[str, Any]:
    from agent_system.agents.documents import summarize_artifact

    target = jailed(str(args.get("path") or ""), ctx.settings, must_exist=True)
    if not target.is_file():
        raise ToolError("document_inspect expects a file")
    try:
        summary = summarize_artifact(target)
    except Exception as exc:
        raise ToolError(f"cannot inspect {target.name}: {exc}") from exc
    return {"summary": summary, "size_bytes": target.stat().st_size}


def _document_extract(args: dict[str, Any], ctx: ToolContext) -> dict[str, Any]:
    target = jailed(str(args.get("path") or ""), ctx.settings, must_exist=True)
    suffix = target.suffix.lower()
    if suffix in (".txt", ".md", ".csv", ".json", ".yaml", ".yml"):
        return {
            "path": relative_label(target),
            "format": suffix,
            "text": read_text(target, ctx)[:MAX_TEXT_CHARS],
        }
    if suffix == ".docx":
        return {"path": relative_label(target), "format": "docx", "text": _docx_text(target)}
    if suffix == ".xlsx":
        return {"path": relative_label(target), "format": "xlsx", "sheets": _xlsx_text(target)}
    if suffix == ".pdf":
        raise ToolError(
            "PDF text extraction is not implemented (no PDF extractor in the dependency set); "
            "read the source content instead"
        )
    raise ToolError(f"document_extract does not support '{suffix or target.name}'")


def _docx_text(path: Path) -> str:
    try:
        import re

        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        text = re.sub(r"<w:p[ >]", "\n<w:p ", xml)
        return re.sub(r"<[^>]+>", "", text).strip()[:MAX_TEXT_CHARS]
    except Exception as exc:
        raise ToolError(f"cannot read docx: {exc}") from exc


def _xlsx_text(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        book = load_workbook(path, read_only=True, data_only=True)
        sheets: list[dict[str, Any]] = []
        for sheet in book.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            sheets.append({"name": sheet.title, "rows": rows[:200], "row_count": len(rows)})
        book.close()
        return sheets
    except Exception as exc:
        raise ToolError(f"cannot read xlsx: {exc}") from exc


def _document_convert(args: dict[str, Any], ctx: ToolContext) -> dict[str, Any]:
    source = jailed(str(args.get("path") or ""), ctx.settings, must_exist=True)
    kind = _kind(args)
    if source.suffix.lower() not in (".txt", ".md", ".csv"):
        raise ToolError(
            "document_convert converts text/markdown/csv sources into pptx/docx/xlsx/pdf/txt/md"
        )
    text = read_text(source, ctx)
    paragraphs = [block.strip() for block in text.split("\n\n") if block.strip()]
    title = source.stem
    if kind in ("txt", "md"):
        content: dict[str, Any] = {"body": text}
    elif kind == "pdf":
        content = {"title": title, "sections": [{"heading": "", "body": p} for p in paragraphs]}
    elif kind == "docx":
        content = {"title": title, "paragraphs": paragraphs}
    elif kind == "pptx":
        content = {
            "title": title,
            "slides": [
                {"title": f"Part {index + 1}", "bullets": [para[:400]]}
                for index, para in enumerate(paragraphs)
            ],
        }
    else:  # xlsx
        rows = [line.split(",") for line in text.splitlines() if line.strip()]
        header = rows[0] if rows else []
        content = {"sheets": [{"name": title[:30], "headers": header, "rows": rows[1:]}]}
    problems = _validate_content(kind, content)
    if problems:
        raise ToolError("conversion produced invalid content: " + "; ".join(problems))
    name = _target_name({**args, "name": str(args.get("name") or source.stem)}, kind)
    if kind in ("txt", "md"):
        target = _outputs_dir(ctx) / name
        target.write_text(text, encoding="utf-8")
    else:
        try:
            target = _document_agent(ctx).generate(kind, name, content)
        except Exception as exc:
            raise ToolError(f"conversion failed: {exc}") from exc
    return {
        "source": relative_label(source),
        "kind": kind,
        "path": str(target),
        "bytes": target.stat().st_size if target.exists() else 0,
    }


def _document_persist(args: dict[str, Any], ctx: ToolContext) -> dict[str, Any]:
    """Copy a produced file into outputs and register it as an Artifact."""
    source = jailed(str(args.get("path") or ""), ctx.settings, must_exist=True)
    if not source.is_file():
        raise ToolError("document_persist expects a file")
    outputs = _outputs_dir(ctx)
    name = str(args.get("name") or source.name)
    if "/" in name or "\\" in name:
        raise ToolError("document_persist: 'name' must be a plain file name")
    target = (outputs / name).resolve()
    if not target.is_relative_to(outputs.resolve()):
        raise ToolError("document_persist: refusing to write outside outputs")
    if source.resolve() != target:
        shutil.copy2(source, target)
    artifact_id = None
    if ctx.factory is not None:
        from agent_system.domain import ids
        from agent_system.infra.db import session_scope
        from agent_system.infra.models import Artifact

        artifact_id = ids.new_id("art")
        with session_scope(ctx.factory) as db:
            db.add(
                Artifact(
                    id=artifact_id,
                    task_id=ctx.task_id,
                    kind=target.suffix.lstrip(".") or "file",
                    path=str(target),
                    size_bytes=target.stat().st_size,
                )
            )
        _emit_artifact(ctx, artifact_id, target)
    return {
        "artifact_id": artifact_id,
        "path": str(target),
        "bytes": target.stat().st_size,
        "recorded": artifact_id is not None,
    }


def _emit_artifact(ctx: ToolContext, artifact_id: str, path: Path) -> None:
    emit = getattr(ctx, "emit", None)
    if emit is None:
        return
    try:
        emit(
            "artifact.created",
            {
                "artifact_id": artifact_id,
                "kind": path.suffix.lstrip(".") or "file",
                "name": path.name,
                "bytes": path.stat().st_size,
            },
        )
    except Exception:
        pass


_CONTENT_SCHEMA = {
    "type": "object",
    "description": (
        "Builder content. pptx: {title, slides[]}; docx: {title, paragraphs[]}; "
        "xlsx: {sheets[{name, headers, rows}]}; pdf: {title, sections[]}; "
        "txt/md: {body}"
    ),
}


def register(registry: ToolRegistry, settings: Any = None) -> None:  # noqa: ARG001
    registry.register(
        Tool(
            name="document_validate",
            description="Validate document content for a kind without writing anything.",
            parameters={
                "type": "object",
                "properties": {
                    "kind": {"type": "string", "enum": list(KINDS)},
                    "content": _CONTENT_SCHEMA,
                },
                "required": ["kind", "content"],
                "additionalProperties": False,
            },
            risk="read",
            handler=_document_validate,
            group=GROUP,
        )
    )
    registry.register(
        Tool(
            name="document_create",
            description="Generate a pptx/docx/xlsx/pdf/txt/md document into outputs.",
            parameters={
                "type": "object",
                "properties": {
                    "kind": {"type": "string", "enum": list(KINDS)},
                    "name": _str_param("File name (extension optional)"),
                    "content": _CONTENT_SCHEMA,
                },
                "required": ["kind", "content"],
                "additionalProperties": False,
            },
            risk="write",
            handler=_document_create,
            scope="document:write",
            group=GROUP,
        )
    )
    registry.register(
        Tool(
            name="document_inspect",
            description="Summarize an existing document artifact (kind, size, structure).",
            parameters={
                "type": "object",
                "properties": {"path": _str_param("Document path")},
                "required": ["path"],
                "additionalProperties": False,
            },
            risk="read",
            handler=_document_inspect,
            group=GROUP,
        )
    )
    registry.register(
        Tool(
            name="document_extract",
            description="Extract text from txt/md/csv/json/docx/xlsx (pdf unsupported).",
            parameters={
                "type": "object",
                "properties": {"path": _str_param("Document path")},
                "required": ["path"],
                "additionalProperties": False,
            },
            risk="read",
            handler=_document_extract,
            group=GROUP,
        )
    )
    registry.register(
        Tool(
            name="document_convert",
            description="Convert text/markdown/csv into another document kind.",
            parameters={
                "type": "object",
                "properties": {
                    "path": _str_param("Source text/markdown/csv file"),
                    "kind": {"type": "string", "enum": list(KINDS)},
                    "name": _str_param("Output file name (default: source stem)"),
                },
                "required": ["path", "kind"],
                "additionalProperties": False,
            },
            risk="write",
            handler=_document_convert,
            scope="document:write",
            group=GROUP,
        )
    )
    registry.register(
        Tool(
            name="document_persist",
            description="Copy a document into outputs and register it as a persisted artifact.",
            parameters={
                "type": "object",
                "properties": {
                    "path": _str_param("File to persist"),
                    "name": _str_param("Output name (default: same file name)"),
                },
                "required": ["path"],
                "additionalProperties": False,
            },
            risk="write",
            handler=_document_persist,
            scope="document:write",
            group=GROUP,
        )
    )


__all__ = ["GROUP", "KINDS", "REQUIRED_CONTENT", "register"]

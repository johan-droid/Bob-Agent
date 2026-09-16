"""DocumentAgent (v3.1 Phase 7): PPTX/DOCX/XLSX/PDF generation.

Deterministic builders — content comes from the task input, so every
generated file is real and openable. Artifacts are written to the outputs
directory; `artifact.created` events are emitted by the caller.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from agent_system.services.secrets import redact_dict

SUPPORTED_KINDS = {"pptx", "docx", "xlsx", "pdf"}


class DocumentError(ValueError):
    pass


class DocumentAgent:
    def __init__(self, outputs_dir: Path) -> None:
        self._dir = outputs_dir
        self._dir.mkdir(parents=True, exist_ok=True)

    def generate(self, kind: str, name: str, content: dict[str, Any]) -> Path:
        kind = kind.lower().lstrip(".")
        if kind not in SUPPORTED_KINDS:
            raise DocumentError(f"unsupported document kind: {kind}")
        safe_name = "".join(c for c in name if c.isalnum() or c in "-_ ").strip() or "document"
        out = self._dir / f"{safe_name}.{kind}"
        builder = {
            "pptx": self._pptx,
            "docx": self._docx,
            "xlsx": self._xlsx,
            "pdf": self._pdf,
        }[kind]
        builder(content, out)
        return out

    # -- builders ------------------------------------------------------------

    def _pptx(self, content: dict[str, Any], out: Path) -> None:
        from pptx import Presentation

        prs = Presentation()
        title = str(content.get("title", "Presentation"))
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = str(content.get("subtitle", ""))
        for section in content.get("slides", []):
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = str(section.get("heading", ""))[:200]
            body = "\n".join(str(b) for b in section.get("bullets", []))
            s.placeholders[1].text = body[:2000]
        prs.save(str(out))

    def _docx(self, content: dict[str, Any], out: Path) -> None:
        from docx import Document

        doc = Document()
        doc.add_heading(str(content.get("title", "Document")), level=0)
        for para in content.get("paragraphs", []):
            doc.add_paragraph(str(para))
        for heading, section in content.get("sections", {}).items():
            doc.add_heading(str(heading), level=1)
            doc.add_paragraph(str(section))
        doc.save(str(out))

    def _xlsx(self, content: dict[str, Any], out: Path) -> None:
        import openpyxl  # type: ignore[import-untyped]

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        sheets = content.get("sheets", {})
        if not sheets:
            raise DocumentError("xlsx content requires 'sheets': {name: [[row], ...]}")
        for sheet_name, rows in sheets.items():
            ws = wb.create_sheet(title=str(sheet_name)[:31])
            for row in rows:
                ws.append(list(row))
        wb.save(str(out))

    def _pdf(self, content: dict[str, Any], out: Path) -> None:
        from fpdf import FPDF
        from fpdf.enums import XPos, YPos

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=16)
        title = str(content.get("title", "Report"))
        # fpdf2 latin-1 core fonts: replace common unicode punctuation.
        pdf.multi_cell(
            0,
            10,
            title.encode("latin-1", "replace").decode("latin-1"),
            new_x=XPos.LMARGIN,
            new_y=YPos.NEXT,
        )
        pdf.set_font("Helvetica", size=11)
        for para in content.get("paragraphs", []):
            pdf.multi_cell(
                0,
                6,
                str(para).encode("latin-1", "replace").decode("latin-1"),
                new_x=XPos.LMARGIN,
                new_y=YPos.NEXT,
            )
            pdf.ln(2)
        pdf.output(str(out))


def summarize_artifact(path: Path) -> dict[str, Any]:
    """Metadata recorded with artifact.created (no content, no secrets)."""
    return redact_dict(
        {
            "path": str(path),
            "kind": path.suffix.lstrip("."),
            "size_bytes": path.stat().st_size,
        }
    )
